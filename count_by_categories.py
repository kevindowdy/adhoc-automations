"""Build a hierarchical row-count table from any ordered list of columns.

Generalizes the top-level / indented-sub-level rollup used in the
vulnerability tracking report (one grouping column nested under another,
each row showing how many raw rows match it) to any dataset and any
number of grouping columns, provided dynamically instead of hardcoded.

Each grouping column becomes one level of the table: a row is emitted for
every distinct value at that level, indented under its parent, with a
count of the raw rows that match the values seen so far. A single
"Grand Total" row closes out the table. No filtering, date parsing, or
categorization logic lives here -- callers are expected to filter the
dataset before it reaches this script.
"""

from __future__ import annotations

import argparse
import logging
from collections.abc import Sequence
from pathlib import Path
from typing import Final

import pandas as pd

# =============================================================================
# CONFIGURATION
# =============================================================================

# Directory used to build a default output path when --output is not given.
DEFAULT_OUTPUT_DIR: Final[Path] = Path("data/output")

# Placeholder substituted for missing/blank values in a grouping column so
# they still get their own counted row instead of being silently dropped.
DEFAULT_BLANK_LABEL: Final[str] = "(Blank)"

# How far each nested level is indented in the rendered "Category" text.
INDENT_SPACES_PER_LEVEL: Final[int] = 2

GRAND_TOTAL_LABEL: Final[str] = "Grand Total"

CATEGORY_COLUMN: Final[str] = "Category"
LEVEL_COLUMN: Final[str] = "Level"
COUNT_COLUMN: Final[str] = "Count"

APP_NAME = "count-by-categories"

# =============================================================================
# LOGGING
# =============================================================================

logging.basicConfig(
    level=logging.INFO,
    format=(
        '{"timestamp":"%(asctime)s","level":"%(levelname)s",'
        f'"message":"%(message)s","application":"{APP_NAME}"}}'
    ),
    handlers=[logging.StreamHandler()],
)

logger = logging.getLogger(__name__)


def _normalize_group_column(series: pd.Series, blank_label: str) -> pd.Series:
    """Coerce a grouping column to trimmed strings with blanks labeled.

    Args:
        series: The column to normalize.
        blank_label: Value substituted for missing or empty-string entries,
            so they group into a single visible row instead of being
            dropped by ``groupby``/``unique``.

    Returns:
        A new string Series with no NaN or empty-string values.
    """
    return series.fillna(blank_label).astype(str).str.strip().replace("", blank_label)


def _build_rows(
    dataframe: pd.DataFrame, remaining_columns: Sequence[str], level: int
) -> list[dict[str, object]]:
    """Recursively build indented count rows for the remaining columns.

    Args:
        dataframe: The (already column-normalized) subset of rows to
            break down further.
        remaining_columns: Grouping columns still to be applied, in
            top-to-bottom nesting order.
        level: Current nesting depth (0 = top level), used for indentation.

    Returns:
        A list of row dicts, each with ``Level``, ``Category``, and
        ``Count`` keys, in the order they should appear in the report
        (each row's descendants immediately follow it).
    """
    if not remaining_columns:
        return []

    column, deeper_columns = remaining_columns[0], remaining_columns[1:]
    indent = " " * (INDENT_SPACES_PER_LEVEL * level)

    rows: list[dict[str, object]] = []
    for value in sorted(dataframe[column].unique()):
        value_df = dataframe[dataframe[column] == value]
        rows.append(
            {
                LEVEL_COLUMN: level,
                CATEGORY_COLUMN: f"{indent}{value}",
                COUNT_COLUMN: len(value_df),
            }
        )
        rows.extend(_build_rows(value_df, deeper_columns, level + 1))

    return rows


def count_by_categories(
    dataframe: pd.DataFrame,
    group_columns: Sequence[str],
    blank_label: str = DEFAULT_BLANK_LABEL,
) -> pd.DataFrame:
    """Build a hierarchical raw row-count table for a dataset.

    One row is produced per distinct value at each grouping level, nested
    directly under its parent row, followed by a trailing "Grand Total"
    row. Counts are raw row counts -- no filtering or other business logic
    is applied here; that is expected to happen before the dataset reaches
    this function.

    Args:
        dataframe: The dataset to summarize. Not mutated.
        group_columns: One or more column names, in top-to-bottom nesting
            order. A single column produces a flat one-level table; two or
            more nest each subsequent column's rows under the matching row
            above it.
        blank_label: Value substituted for missing/blank entries in any
            grouping column.

    Returns:
        A DataFrame with columns ``Level``, ``Category``, and ``Count``,
        one row per group value (indented under its parent by
        ``INDENT_SPACES_PER_LEVEL`` spaces per level in ``Category``) plus
        a final Grand Total row.

    Raises:
        ValueError: If ``group_columns`` is empty or contains duplicates.
        KeyError: If any column in ``group_columns`` is not present in
            ``dataframe``.
    """
    if not group_columns:
        raise ValueError("At least one group-by column is required.")

    if len(set(group_columns)) != len(group_columns):
        raise ValueError(
            f"group_columns must not contain duplicates: {group_columns!r}"
        )

    missing_columns = [col for col in group_columns if col not in dataframe.columns]
    if missing_columns:
        raise KeyError(
            f"Column(s) not found in dataset: {missing_columns}. "
            f"Available columns: {list(dataframe.columns)}"
        )

    working = dataframe.copy()
    for column in group_columns:
        working[column] = _normalize_group_column(working[column], blank_label)

    rows = _build_rows(working, group_columns, level=0)
    rows.append(
        {
            LEVEL_COLUMN: "",
            CATEGORY_COLUMN: GRAND_TOTAL_LABEL,
            COUNT_COLUMN: len(working),
        }
    )

    return pd.DataFrame(rows, columns=[LEVEL_COLUMN, CATEGORY_COLUMN, COUNT_COLUMN])


def _read_dataset(path: Path, sheet_name: str | int | None) -> pd.DataFrame:
    """Read a CSV or Excel dataset into a string-typed DataFrame.

    Args:
        path: Path to the input file. Must exist.
        sheet_name: Sheet name or zero-based index to read when ``path``
            is an Excel workbook. Ignored for CSV/TSV input. Defaults to
            the first sheet when ``None``.

    Returns:
        The loaded dataset, with every column read as string dtype so
        grouping treats equal-looking values consistently.

    Raises:
        FileNotFoundError: If ``path`` does not exist.
        ValueError: If ``path`` has an unsupported extension.
    """
    if not path.is_file():
        raise FileNotFoundError(f"Input file not found: {path}")

    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path, dtype=str)
    if suffix == ".tsv":
        return pd.read_csv(path, dtype=str, sep="\t")
    if suffix in (".xlsx", ".xls"):
        return pd.read_excel(
            path, dtype=str, sheet_name=sheet_name if sheet_name is not None else 0
        )

    raise ValueError(
        f"Unsupported input file extension {suffix!r} for {path}. "
        "Supported extensions: .csv, .tsv, .xlsx, .xls"
    )


def _style_workbook(worksheet, report_df: pd.DataFrame) -> None:
    """Apply header, indentation, and Grand Total styling to a report sheet.

    Args:
        worksheet: The openpyxl worksheet the report was just written to.
        report_df: The report DataFrame that was written, used to look up
            each row's nesting level for styling.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    DARK_BLUE = "1F4E79"
    LIGHT_BLUE = "D6E4F0"
    WHITE = "FFFFFF"

    for cell in worksheet[1]:
        cell.fill = PatternFill(
            start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid"
        )
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    worksheet.row_dimensions[1].height = 32

    for row_idx, level in enumerate(report_df[LEVEL_COLUMN].tolist(), start=2):
        is_grand_total = level == ""
        is_top_level = level == 0

        for col_idx, cell in enumerate(worksheet[row_idx], start=1):
            if is_grand_total:
                cell.fill = PatternFill(
                    start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid"
                )
                cell.font = Font(color=WHITE, bold=True)
                cell.alignment = Alignment(horizontal="center")
            elif is_top_level:
                cell.fill = PatternFill(
                    start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type="solid"
                )
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
                if col_idx == 2:
                    cell.alignment = Alignment(horizontal="left")
            else:
                cell.alignment = Alignment(horizontal="center")
                if col_idx == 2:
                    cell.alignment = Alignment(horizontal="left", indent=int(level))

    for col_idx, col_cells in enumerate(worksheet.columns, start=1):
        max_len = max((len(str(cell.value or "")) for cell in col_cells), default=0)
        col_letter = get_column_letter(col_idx)
        worksheet.column_dimensions[col_letter].width = (
            min(max_len + 3, 60) if col_idx == 2 else max(9, min(max_len + 2, 18))
        )

    worksheet.freeze_panes = "A2"


def save_report(report_df: pd.DataFrame, output_path: Path) -> None:
    """Write a report DataFrame to CSV or styled Excel, atomically.

    The file is staged to a temporary path in the same directory and only
    moved into place once the write completes successfully, so a failed
    or interrupted write never leaves a partial file at ``output_path``.

    Args:
        report_df: The report to write, as returned by
            :func:`count_by_categories`.
        output_path: Destination path. Extension (``.csv`` or ``.xlsx``)
            selects the output format.

    Raises:
        ValueError: If ``output_path`` has an unsupported extension.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    suffix = output_path.suffix.lower()
    temp_output_path = output_path.with_suffix(f"{output_path.suffix}.tmp")

    try:
        if suffix == ".csv":
            report_df.to_csv(temp_output_path, index=False)
        elif suffix == ".xlsx":
            with pd.ExcelWriter(temp_output_path, engine="openpyxl") as writer:
                report_df.to_excel(writer, index=False, sheet_name="Category Counts")
                _style_workbook(writer.sheets["Category Counts"], report_df)
        else:
            raise ValueError(
                f"Unsupported output file extension {suffix!r} for {output_path}. "
                "Supported extensions: .csv, .xlsx"
            )

        temp_output_path.replace(output_path)
    finally:
        temp_output_path.unlink(missing_ok=True)


def _parse_sheet_name(value: str) -> str | int:
    """Parse a --sheet-name CLI value as a sheet index when numeric.

    Args:
        value: Raw CLI argument string.

    Returns:
        ``int(value)`` if it looks like a plain integer, otherwise the
        string unchanged (treated as a sheet name by pandas).
    """
    return int(value) if value.lstrip("-").isdigit() else value


def _default_output_path(input_file: Path) -> Path:
    """Build a default output path from the input file's name.

    Args:
        input_file: The dataset being summarized.

    Returns:
        ``DEFAULT_OUTPUT_DIR / "<input stem>_category_counts.xlsx"``.
    """
    return DEFAULT_OUTPUT_DIR / f"{input_file.stem}_category_counts.xlsx"


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    """Parse command-line arguments for the script.

    Args:
        argv: Argument list to parse, or ``None`` to use ``sys.argv``.

    Returns:
        The parsed arguments.
    """
    parser = argparse.ArgumentParser(
        description=(
            "Count raw rows in a dataset, grouped by any ordered list of "
            "columns, with each column nested under the one before it."
        )
    )
    parser.add_argument(
        "input_file", type=Path, help="Path to a CSV, TSV, or Excel dataset."
    )
    parser.add_argument(
        "--columns",
        "-c",
        nargs="+",
        required=True,
        metavar="COLUMN",
        help="Grouping columns, top-to-bottom nesting order (e.g. --columns Owner App).",
    )
    parser.add_argument(
        "--output",
        "-o",
        type=Path,
        default=None,
        help=f"Output path (.csv or .xlsx). Defaults to {DEFAULT_OUTPUT_DIR}/<input stem>_category_counts.xlsx",
    )
    parser.add_argument(
        "--sheet-name",
        type=_parse_sheet_name,
        default=None,
        help="Sheet name or 0-based index to read for Excel input (defaults to the first sheet).",
    )
    parser.add_argument(
        "--blank-label",
        default=DEFAULT_BLANK_LABEL,
        help=f"Label used for missing/blank grouping values (default: {DEFAULT_BLANK_LABEL!r}).",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> None:
    """Load a dataset, build the hierarchical count report, and save it.

    Args:
        argv: Argument list to parse, or ``None`` to use ``sys.argv``.
    """
    args = parse_args(argv)
    output_path = args.output or _default_output_path(args.input_file)

    logger.info("=" * 80)
    logger.info("Count-By-Categories SCRIPT STARTED")
    logger.info(f"Input file: {args.input_file}")
    logger.info(f"Group columns: {args.columns}")
    logger.info(f"Output file: {output_path}")

    dataframe = _read_dataset(args.input_file, args.sheet_name)
    logger.info(f"  {len(dataframe):,} rows loaded")

    report_df = count_by_categories(dataframe, args.columns, args.blank_label)
    logger.info(
        f"Report built: {len(report_df):,} rows across {len(args.columns)} level(s)"
    )

    save_report(report_df, output_path)
    logger.info(f"Report saved to {output_path}")

    logger.info("SCRIPT COMPLETED SUCCESSFULLY")


if __name__ == "__main__":
    main()
